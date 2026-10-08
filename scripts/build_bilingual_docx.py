from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import platform
import re
import sys

if len(sys.argv) != 3:
    raise SystemExit("Usage: python build_bilingual_docx.py <lesson-content.txt> <output.docx>")

CONTENT = Path(sys.argv[1])
DOCX = Path(sys.argv[2])
SUPERSCRIPT = str.maketrans("0123456789", "⁰¹²³⁴⁵⁶⁷⁸⁹")
CJK_FONT = {"Windows": "Microsoft YaHei", "Darwin": "PingFang SC"}.get(platform.system(), "Noto Sans CJK SC")


def normalize_math(text):
    text = text.replace("sqrt(1^2 + (-sqrt(3))^2)", "√(1² + (-√3)²)")
    text = text.replace("sqrt((-2)^2 + 2^2)", "√((-2)² + 2²)")
    for old, new in (("+/-", "±"), ("->", "→"), ("Delta", "Δ")):
        text = text.replace(old, new)
    for _ in range(5):
        changed = re.sub(r"sqrt\(([^()]*)\)", r"√(\1)", text)
        if changed == text:
            break
        text = changed
    text = re.sub(r"abs\(([^()]*)\)", r"|\1|", text)
    text = re.sub(r"\bpi\b", "π", text)
    text = re.sub(r"(\d)pi", r"\1π", text)
    text = re.sub(r"\balpha\b", "α", text)
    text = re.sub(r"\bbeta\b", "β", text)
    text = re.sub(r"\btheta\b", "θ", text)
    text = re.sub(r"([A-Za-z0-9)])\^(\d+)", lambda m: m.group(1) + m.group(2).translate(SUPERSCRIPT), text)
    text = re.sub(r"\^\((\d+)\)", lambda m: m.group(1).translate(SUPERSCRIPT), text)
    text = re.sub(r"(?<=\S)\s+\*\s+(?=\S)", " × ", text)
    text = re.sub(r"(?<=\d)\s*\*\s*(?=\d)", " × ", text)
    return text

doc = Document()
for section in doc.sections:
    section.top_margin = Inches(0.65); section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75); section.right_margin = Inches(0.75)

normal = doc.styles["Normal"]
normal.font.name = "Calibri"; normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
normal.paragraph_format.space_after = Pt(4); normal.paragraph_format.line_spacing = 1.08

def set_style_font(style, name="Calibri", east=CJK_FONT, size=10.5, bold=None, color=None):
    style.font.name = name; style.font.size = Pt(size)
    if bold is not None: style.font.bold = bold
    if color is not None: style.font.color.rgb = RGBColor.from_string(color)
    rpr = style.element.get_or_add_rPr(); rf = rpr.rFonts
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:ascii"), name); rf.set(qn("w:hAnsi"), name); rf.set(qn("w:eastAsia"), east)

styles = doc.styles
for name in ["EnglishLine","ChineseLine","FormulaLine","ChineseHeading","QuestionEnglish","QuestionChinese","TinyNote"]:
    if name not in styles: styles.add_style(name, 1)
set_style_font(styles["EnglishLine"], size=11, bold=True, color="17365D")
styles["EnglishLine"].paragraph_format.space_before = Pt(2); styles["EnglishLine"].paragraph_format.space_after = Pt(1); styles["EnglishLine"].paragraph_format.keep_with_next = True
set_style_font(styles["ChineseLine"], size=10, color="3F3F3F")
styles["ChineseLine"].paragraph_format.space_after = Pt(6); styles["ChineseLine"].paragraph_format.keep_with_next = True
set_style_font(styles["FormulaLine"], name="Cambria Math", east=CJK_FONT, size=10.5, color="111111")
styles["FormulaLine"].paragraph_format.left_indent = Inches(0.22); styles["FormulaLine"].paragraph_format.space_after = Pt(1); styles["FormulaLine"].paragraph_format.keep_with_next = True
set_style_font(styles["ChineseHeading"], size=10, bold=True, color="595959")
styles["ChineseHeading"].paragraph_format.space_after = Pt(6); styles["ChineseHeading"].paragraph_format.keep_with_next = True
set_style_font(styles["QuestionEnglish"], size=10.5, bold=True, color="000000")
styles["QuestionEnglish"].paragraph_format.left_indent = Inches(0.15); styles["QuestionEnglish"].paragraph_format.space_after = Pt(1); styles["QuestionEnglish"].paragraph_format.keep_with_next = True
set_style_font(styles["QuestionChinese"], size=9.5, color="404040")
styles["QuestionChinese"].paragraph_format.left_indent = Inches(0.15); styles["QuestionChinese"].paragraph_format.space_after = Pt(5); styles["QuestionChinese"].paragraph_format.keep_with_next = True
set_style_font(styles["TinyNote"], size=8.5, color="666666")
for hname, size, before, after in [("Title",23,0,6),("Heading 1",16,12,2),("Heading 2",12.5,8,2),("Heading 3",11.2,6,2)]:
    st = styles[hname]; set_style_font(st, size=size, bold=True, color="17365D")
    st.paragraph_format.keep_with_next = True; st.paragraph_format.space_before = Pt(before); st.paragraph_format.space_after = Pt(after)

def add_page_field(paragraph):
    run = paragraph.add_run()
    f1 = OxmlElement("w:fldChar"); f1.set(qn("w:fldCharType"), "begin")
    ins = OxmlElement("w:instrText"); ins.set(qn("xml:space"), "preserve"); ins.text = " PAGE "
    f2 = OxmlElement("w:fldChar"); f2.set(qn("w:fldCharType"), "end")
    run._r.append(f1); run._r.append(ins); run._r.append(f2)

footer = doc.sections[0].footer; fp = footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = fp.add_run("Bilingual Teaching Plan | "); r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string("777777")
add_page_field(fp)
for run in fp.runs:
    run.font.name = "Calibri"; run.font.size = Pt(8); run.font.color.rgb = RGBColor.from_string("777777")

def shade_p(p, fill):
    pPr = p._p.get_or_add_pPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); pPr.append(shd)

def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement("w:shd"); shd.set(qn("w:fill"), fill); tcPr.append(shd)

def page_break(): doc.add_page_break()

def fmt_run(run, name="Calibri", east=CJK_FONT, size=10.5, bold=False, color="000000"):
    run.font.name = name; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr(); rf = rpr.rFonts
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:ascii"), name); rf.set(qn("w:hAnsi"), name); rf.set(qn("w:eastAsia"), east)

def pair_para(en, zh, kind="p"):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before = Pt(2 if kind == "p" else 1); pf.space_after = Pt(6 if kind == "p" else 4)
    pf.keep_together = True; pf.keep_with_next = True
    if kind == "bullet": pf.left_indent = Inches(0.24)
    if kind == "formula": pf.left_indent = Inches(0.22)
    if kind == "question": pf.left_indent = Inches(0.15)
    r1 = p.add_run(("• " if kind == "bullet" else "") + normalize_math(en))
    fmt_run(r1, name=("Cambria Math" if kind == "formula" else "Calibri"), size=(11 if kind in ("p","bullet") else 10.5), bold=(kind in ("p","bullet","question")), color=("17365D" if kind in ("p","bullet") else "000000"))
    br = p.add_run(); br.add_break(WD_BREAK.LINE)
    r2 = p.add_run(("• " if kind == "bullet" else "") + normalize_math(zh))
    fmt_run(r2, size=(10 if kind in ("p","bullet","formula") else 9.5), bold=False, color=("3F3F3F" if kind in ("p","bullet","formula") else "404040"))
    return p

def add_pair(en, zh): return pair_para(en, zh, "p")
def add_bullet(en, zh): return pair_para(en, zh, "bullet")
def add_formula(en, zh): return pair_para(en, zh, "formula")
def add_question(num, en, zh): return pair_para(f"{num}. {en}", f"{num}. {zh}", "question")

def heading_pair(en, zh, style_name, zh_size):
    p = doc.add_paragraph(style=style_name)
    p.paragraph_format.keep_together = True
    p.paragraph_format.keep_with_next = True
    r1 = p.add_run(normalize_math(en))
    br = p.add_run(); br.add_break(WD_BREAK.LINE)
    r2 = p.add_run(normalize_math(zh))
    fmt_run(r2, name=CJK_FONT, size=zh_size, bold=True, color="595959")
    return p

def set_cell_text(cell, text):
    pieces = normalize_math(text).split("[[BR]]")
    cell.text = ""
    p = cell.paragraphs[0]
    p.add_run(pieces[0])
    for piece in pieces[1:]:
        p.add_run().add_break(WD_BREAK.LINE)
        p.add_run(piece)
    p.paragraph_format.keep_together = True

def add_table(header, rows, widths):
    table = doc.add_table(rows=1, cols=len(header)); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]; set_cell_text(cell, text); shade_cell(cell, "DCE6F1"); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in cell.paragraphs[0].runs: r.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row): set_cell_text(cells[i], text); cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    old_bold = run.bold
                    fmt_run(run, name="Calibri", size=9.5, bold=bool(old_bold), color="000000")
    note = doc.add_paragraph(style="TinyNote")
    note.add_run("English first, Chinese second in each cell.")
    note.add_run().add_break(WD_BREAK.LINE)
    note.add_run("每格先英文，后中文。")

lines = CONTENT.read_text(encoding="utf-8-sig").splitlines(); i = 0
while i < len(lines):
    line = lines[i].strip(); i += 1
    if not line or line.startswith("#"): continue
    parts = line.split("||"); kind = parts[0]
    if kind == "page": page_break()
    elif kind == "h1": heading_pair(parts[1], parts[2], "Heading 1", 10)
    elif kind == "h2": heading_pair(parts[1], parts[2], "Heading 2", 10)
    elif kind == "h3": heading_pair(parts[1], parts[2], "Heading 3", 9.5)
    elif kind == "p": add_pair(parts[1], parts[2])
    elif kind == "b": add_bullet(parts[1], parts[2])
    elif kind == "f": add_formula(parts[1], parts[2])
    elif kind == "q": add_question(parts[1], parts[2], parts[3])
    elif kind == "title":
        p = doc.add_paragraph(style="Title"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(parts[1])
        p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER; rr = p2.add_run(parts[2]); rr.bold = True; fmt_run(rr, name=CJK_FONT, size=14, bold=True, color="17365D")
    elif kind == "subtitle":
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; rr = p.add_run(parts[1]); fmt_run(rr, size=11, color="595959")
        p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER; rr = p2.add_run(parts[2]); fmt_run(rr, name=CJK_FONT, size=11, color="595959")
    elif kind == "table":
        widths = [float(x) for x in parts[2].split(",")]
        header = lines[i].split("||")[1:]; i += 1; rows = []
        while i < len(lines) and lines[i].startswith("row||"): rows.append(lines[i].split("||")[1:]); i += 1
        add_table(header, rows, widths)

DOCX.parent.mkdir(parents=True, exist_ok=True)
doc.save(str(DOCX)); print(DOCX)
